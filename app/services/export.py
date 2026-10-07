"""Excel export of a control assessment: the workbook an auditor gets."""
import io
import json

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models import Artifact, Assessment, AuditEvent, EvidenceLink, Issue, User, utcnow
from app.services import assessments, responses

RATING = responses.RATINGS


def _cell(value):
    if value is None:
        return ""
    if hasattr(value, "isoformat"):
        return value.isoformat(sep=" ", timespec="minutes") if hasattr(value, "hour") else value.isoformat()
    if isinstance(value, (dict, list)):
        value = json.dumps(value, default=str)
    if isinstance(value, str):
        return ILLEGAL_CHARACTERS_RE.sub("", value)
    return value


def _sheet(wb: Workbook, title: str, header: list[str], rows: list[list], widths: dict | None = None):
    ws = wb.create_sheet(title)
    ws.append(header)
    for c in ws[1]:
        c.font = Font(bold=True)
    for row in rows:
        ws.append([_cell(v) for v in row])
        for c in ws[ws.max_row]:
            if isinstance(c.value, str):
                if c.value.startswith("="):
                    c.data_type = "s"  # user text, never a formula
                if "\n" in c.value:
                    c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for col, width in (widths or {}).items():
        ws.column_dimensions[col].width = width
    return ws


def _links(links: list[EvidenceLink]) -> str:
    out = []
    for link in links:
        if link.kind == "artifact" and link.artifact:
            out.append(f"{link.title or link.artifact.filename} [file sha256 {link.artifact.sha256}]")
        else:
            pin = "" if link.pinned is None else (" [pinned]" if link.pinned else " [NOT pinned]")
            out.append(f"{link.title + ': ' if link.title else ''}{link.url}{pin}")
    return "\n".join(out)


def controls_workbook(db: Session, a: Assessment, user: User) -> bytes:
    tree = assessments.Tree(a.scope_nodes)
    by_item = assessments.responses_by_item(db, a.id)
    rollup = assessments.rollup(db, a)
    live_links: dict[int, list[EvidenceLink]] = {}
    for link in db.scalars(select(EvidenceLink).where(EvidenceLink.assessment_id == a.id,
                                                      EvidenceLink.removed_at.is_(None))):
        live_links.setdefault(link.response_id, []).append(link)

    wb = Workbook()
    wb.remove(wb.active)
    totals: dict[str, int] = {}
    for row in rollup:
        for o, n in row["outcomes"].items():
            totals[o] = totals.get(o, 0) + n
    att = a.attestation or {}
    summary = [
        ["Assessment", a.name], ["Status", a.status],
        ["Period", f"{_cell(a.period_start)} to {_cell(a.period_end)}"],
        ["Scope", tree.root.name], ["Subscriptions", len(tree.leaves())],
        ["Controls library", a.library_snapshot.get("name", "")],
        ["Control document version", a.library_snapshot.get("source_version", "")],
        ["Control document", a.library_snapshot.get("source_url", "")],
        ["Controls", len(a.items)],
        ["Created", f"{_cell(a.created_at)} by {a.created_by}"],
        ["Closed", f"{att.get('closed_at', '')} by {att.get('closed_by', '')}" if att else "not closed"],
        ["Attestation", att.get("statement", "")],
        ["Exported", f"{_cell(utcnow())} by {user.upn}"],
        [], ["Effective outcome (control x subscription)", "Count"],
        *[[k, v] for k, v in sorted(totals.items())],
    ]
    _sheet(wb, "Summary", ["Field", "Value"], summary, {"A": 42, "B": 90})

    _sheet(wb, "Controls", ["Ref", "Title", "Category", "Frameworks", "Owner", "Frequency",
                            "Effective", "Partially effective", "Ineffective", "Not tested", "N/A",
                            "In progress", "Unanswered"],
           [[r["item"].ref, r["item"].title, r["item"].category, r["item"].framework_refs,
             r["item"].owner, r["item"].frequency,
             *[r["outcomes"].get(o, 0) for o in ("effective", "partially_effective", "ineffective",
                                                 "not_tested", "na", "in_progress", "unanswered")]]
            for r in rollup], {"B": 50, "D": 30})

    eff_rows = []
    for item in a.items:
        own = by_item.get(item.id, {})
        for leaf in tree.leaves():
            r, src = assessments.effective(tree, own, leaf.id)
            answered_at = (src.name if src else leaf.name) if r else ""
            eff_rows.append([
                item.ref, item.title, tree.path(leaf.id), leaf.name, leaf.azure_id or "",
                assessments.outcome(a.kind, r), answered_at, "inherited" if src else ("own" if r else ""),
                *(_response_cells(r) if r else [""] * 11),
                _links(live_links.get(r.id, [])) if r else ""])
    _sheet(wb, "Effective answers", [
        "Ref", "Control", "Scope path", "Subscription", "Subscription id", "Outcome", "Answered at",
        "Own / inherited", "Applicability", "Design", "Operating", "How performed", "Test performed",
        "Test result", "Exceptions", "Review state", "Prepared by / at", "Reviewed by / at",
        "Last saved by", "Evidence"], eff_rows, {"B": 40, "C": 40, "L": 60, "M": 40, "N": 40, "T": 60})

    own_rows = []
    for item in a.items:
        for node_id, r in sorted(by_item.get(item.id, {}).items()):
            own_rows.append([item.ref, item.title, tree.path(node_id), tree.nodes[node_id].kind,
                             *_response_cells(r), _links(live_links.get(r.id, []))])
    _sheet(wb, "Answers (all levels)", [
        "Ref", "Control", "Answered at", "Level", "Applicability", "Design", "Operating",
        "How performed", "Test performed", "Test result", "Exceptions", "Review state",
        "Prepared by / at", "Reviewed by / at", "Last saved by", "Evidence"], own_rows,
        {"B": 40, "C": 40, "H": 60, "P": 60})

    issues = db.scalars(select(Issue).where(Issue.assessment_id == a.id).order_by(Issue.id)).all()
    _sheet(wb, "Issues", ["Issue", "Control", "Scope", "Title", "Description", "Severity", "Status",
                          "Owner", "Due", "Remediation plan", "Raised by", "Raised at", "Closed at"],
           [[i.ref, i.item.ref if i.item else "", tree.path(i.scope_node_id) if i.scope_node_id else "",
             i.title, i.description, i.severity, i.status, i.owner, i.due_date, i.remediation_plan,
             i.created_by, i.created_at, i.closed_at] for i in issues], {"D": 40, "E": 50, "J": 50})

    all_links = db.scalars(select(EvidenceLink).where(EvidenceLink.assessment_id == a.id)
                           .order_by(EvidenceLink.id)).all()
    resp_ref = {r.id: (r.item.ref, tree.path(r.scope_node_id))
                for own in by_item.values() for r in own.values()}
    _sheet(wb, "Evidence", ["Control", "Answered at", "Issue", "Kind", "Title", "URL", "File",
                            "sha256", "Pinned to commit", "Added by", "Added at", "Removed by",
                            "Removed at"],
           [[*resp_ref.get(link.response_id, ("", "")), f"ISS-{link.issue_id:04d}" if link.issue_id else "",
             link.kind, link.title, link.url, link.artifact.filename if link.artifact else "",
             link.artifact.sha256 if link.artifact else "",
             "" if link.pinned is None else ("yes" if link.pinned else "NO"),
             link.added_by, link.added_at, link.removed_by, link.removed_at] for link in all_links],
           {"E": 40, "F": 60})

    files = db.scalars(select(Artifact).where(Artifact.assessment_id == a.id).order_by(Artifact.id)).all()
    _sheet(wb, "Files", ["Id", "Filename", "Size (bytes)", "sha256", "Uploaded by", "Uploaded at"],
           [[f.id, f.filename, f.size, f.sha256, f.uploaded_by, f.uploaded_at] for f in files])

    events = db.scalars(select(AuditEvent).where(AuditEvent.assessment_id == a.id)
                        .order_by(AuditEvent.id)).all()
    _sheet(wb, "Audit log", ["#", "At", "Actor", "Action", "Entity", "Entity id", "Changes", "Note"],
           [[e.id, e.at, e.actor, e.action, e.entity_type, e.entity_id, e.changes, e.note]
            for e in events], {"G": 100, "H": 40})

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _response_cells(r) -> list:
    return ["N/A" if r.status == "na" else "Applicable", RATING.get(r.design_rating, r.design_rating),
            RATING.get(r.operating_rating, r.operating_rating), r.narrative, r.test_procedure,
            r.test_result, r.exceptions, r.review_state,
            f"{r.prepared_by} {_cell(r.prepared_at)}".strip(),
            f"{r.reviewed_by} {_cell(r.reviewed_at)}".strip(), r.updated_by]
