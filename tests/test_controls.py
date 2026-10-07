"""Control assessments (RCSA): libraries, scope sync, inheritance, sign-off, close, issues, export."""
import io
from urllib.parse import unquote_plus as unquote

import pytest
from openpyxl import load_workbook
from sqlalchemy import select

from app import seed
from app.models import AuditEvent, Issue, Library, LibraryItem, Response, ScopeNode
from app.services import assessments, scope
from conftest import ADMIN, PREPARER, REVIEWER, VIEWER, aid_from, as_user, item_id, node_id

TSV = ("Control ID\tControl name\tDomain\tFrameworks\tControl owner\n"
       "IAM-01\tMFA for privileged access\tIdentity\tNIST IA-2(1); CIS 1.1\tIAM team\n"
       "LOG-01\tActivity logs retained\tLogging\tNIST AU-11\tPlatform\n")


@pytest.fixture
def setup(client, session):
    """Library imported from a paste, demo scope tree, one assessment over the whole tenant."""
    lid = aid_from(client.post("/libraries", data={"name": "Cloud Control Standard",
                                                   "source_version": "v3.2"}, follow_redirects=False))
    resp = client.post(f"/libraries/{lid}/import", data={"text": TSV}, follow_redirects=False)
    assert "2 added" in unquote(resp.headers["location"])
    seed.load_demo(session)
    session.commit()
    tenant = session.scalar(select(ScopeNode).where(ScopeNode.kind == "tenant"))
    aid = aid_from(client.post("/assessments/controls", data={
        "name": "Q4 RCSA", "library_id": lid, "root_id": tenant.id,
        "period_start": "2026-10-01", "period_end": "2026-12-31"}, follow_redirects=False))
    return {"lid": lid, "aid": aid}


def answer(client, session, aid, ref, node, version=0, **fields):
    return client.post(f"/assessments/{aid}/responses", data={
        "item_id": item_id(session, aid, ref), "node_id": node_id(session, aid, node),
        "version": version, **fields})


COMPLETE = {"narrative": "Conditional Access policy CA-01 enforces MFA.",
            "design_rating": "effective", "operating_rating": "effective",
            "test_procedure": "Reviewed CA policy export", "test_result": "Policy on, no exclusions"}


def transition(client, session, aid, ref, node, action, comment=""):
    return client.post(f"/assessments/{aid}/responses/transition", data={
        "item_id": item_id(session, aid, ref), "node_id": node_id(session, aid, node),
        "action": action, "comment": comment}, follow_redirects=False)


def test_import_parses_excel_paste_and_upserts(client, session, setup):
    lib = session.get(Library, setup["lid"])
    assert [(i.ref, i.category, i.owner) for i in lib.items] == [
        ("IAM-01", "Identity", "IAM team"), ("LOG-01", "Logging", "Platform")]
    resp = client.post(f"/libraries/{setup['lid']}/import", data={
        "text": "ref,title\nIAM-01,MFA for ALL privileged access\nNET-01,No open RDP\n"},
        follow_redirects=False)
    assert "1 added, 1 updated" in unquote(resp.headers["location"])
    resp = client.post(f"/libraries/{setup['lid']}/import", data={"text": "foo,bar\n1,2\n"},
                       follow_redirects=False)
    assert "Nothing imported" in unquote(resp.headers["location"])


def test_assessment_is_a_frozen_snapshot(client, session, setup):
    aid = setup["aid"]
    lib_item = session.scalar(select(LibraryItem).where(LibraryItem.ref == "IAM-01"))
    client.post(f"/libraries/{setup['lid']}/items/{lib_item.id}", data={
        "ref": "IAM-01", "title": "Changed later", "active": "on"})
    page = client.get(f"/assessments/{aid}")
    assert "MFA for privileged access" in page.text and "Changed later" not in page.text
    names = {n.name for n in session.get(assessments.Assessment, aid).scope_nodes}
    assert {"mg-prod", "sub-app-a-prod", "sub-app-a-dev"} <= names


def test_inheritance_and_override(client, session, setup):
    aid = setup["aid"]
    assert answer(client, session, aid, "IAM-01", "mg-prod", **COMPLETE).status_code == 200
    a = session.get(assessments.Assessment, aid)
    tree = assessments.Tree(a.scope_nodes)
    own = assessments.responses_by_item(session, aid)[item_id(session, aid, "IAM-01")]
    sub = node_id(session, aid, "sub-app-b-prod")
    r, src = assessments.effective(tree, own, sub)
    assert src.name == "mg-prod" and r.narrative.startswith("Conditional Access")
    assert assessments.effective(tree, own, node_id(session, aid, "sub-app-a-dev")) == (None, None)

    iid = item_id(session, aid, "IAM-01")
    page = client.get(f"/assessments/{aid}/items/{iid}?node={sub}")
    assert "Inherited from" in page.text and "mg-prod" in page.text

    client.post(f"/assessments/{aid}/items/{iid}/override", data={"node_id": sub, "mode": "copy"})
    session.expire_all()
    copied = session.scalar(select(Response).where(Response.scope_node_id == sub))
    assert copied.narrative == COMPLETE["narrative"] and copied.review_state == "draft"
    rollup = {r["item"].ref: r for r in assessments.rollup(session, a)}["IAM-01"]
    assert rollup["outcomes"] == {"effective": 2, "unanswered": 3}

    # revert: the override is deleted, the deletion is audited, inheritance resumes
    client.post(f"/assessments/{aid}/items/{iid}/inherit", data={"node_id": sub})
    session.expire_all()
    assert session.scalar(select(Response).where(Response.scope_node_id == sub)) is None
    deleted = session.scalar(select(AuditEvent).where(AuditEvent.action == "delete"))
    assert deleted.changes["narrative"][0] == COMPLETE["narrative"]


def test_na_override(client, session, setup):
    aid = setup["aid"]
    answer(client, session, aid, "LOG-01", "Contoso tenant (demo)", **COMPLETE)
    iid = item_id(session, aid, "LOG-01")
    dev = node_id(session, aid, "sub-app-a-dev")
    client.post(f"/assessments/{aid}/items/{iid}/override", data={"node_id": dev, "mode": "na"})
    session.expire_all()
    r = session.scalar(select(Response).where(Response.scope_node_id == dev))
    assert r.status == "na"
    resp = transition(client, session, aid, "LOG-01", "sub-app-a-dev", "prepare")
    assert "narrative" in unquote(resp.headers["location"])  # N/A needs a reason


def test_signoff_flow_and_segregation_of_duties(client, session, setup):
    aid = setup["aid"]
    root = "Contoso tenant (demo)"
    as_user(client, PREPARER)
    answer(client, session, aid, "IAM-01", root, narrative="partial answer")
    resp = transition(client, session, aid, "IAM-01", root, "prepare")
    assert "Fill in" in unquote(resp.headers["location"])
    answer(client, session, aid, "IAM-01", root, version=1, **COMPLETE)
    transition(client, session, aid, "IAM-01", root, "prepare")

    resp = transition(client, session, aid, "IAM-01", root, "review")
    assert resp.status_code == 303 and "requires the reviewer role" in unquote(resp.headers["location"])

    as_user(client, REVIEWER)
    resp = transition(client, session, aid, "IAM-01", root, "return")
    assert "comment is required" in unquote(resp.headers["location"])
    transition(client, session, aid, "IAM-01", root, "return", "add the CA policy export")
    session.expire_all()
    r = session.scalar(select(Response).where(Response.assessment_item_id == item_id(session, aid, "IAM-01")))
    assert r.review_state == "returned"

    as_user(client, PREPARER)
    transition(client, session, aid, "IAM-01", root, "prepare")
    as_user(client, REVIEWER)
    transition(client, session, aid, "IAM-01", root, "review")
    session.expire_all()
    r = session.get(Response, r.id)
    assert (r.review_state, r.prepared_by, r.reviewed_by) == ("reviewed", PREPARER, REVIEWER)

    # signed off -> locked until reopened
    as_user(client, PREPARER)
    assert answer(client, session, aid, "IAM-01", root, version=r.version,
                  narrative="sneaky edit").status_code == 423
    as_user(client, REVIEWER)
    transition(client, session, aid, "IAM-01", root, "reopen", "policy changed in November")
    as_user(client, PREPARER)
    session.expire_all()
    assert answer(client, session, aid, "IAM-01", root, version=session.get(Response, r.id).version,
                  narrative="updated").status_code == 200


def test_reviewer_cannot_sign_off_own_work(client, session, setup):
    aid = setup["aid"]
    root = "Contoso tenant (demo)"
    as_user(client, REVIEWER)
    answer(client, session, aid, "IAM-01", root, **COMPLETE)
    transition(client, session, aid, "IAM-01", root, "prepare")
    resp = transition(client, session, aid, "IAM-01", root, "review")
    assert "Segregation of duties" in unquote(resp.headers["location"])


def test_editing_prepared_answer_returns_it_to_draft(client, session, setup):
    aid = setup["aid"]
    root = "Contoso tenant (demo)"
    answer(client, session, aid, "IAM-01", root, **COMPLETE)
    transition(client, session, aid, "IAM-01", root, "prepare")
    session.expire_all()
    r = session.scalar(select(Response))
    resp = answer(client, session, aid, "IAM-01", root, version=r.version, narrative="changed")
    assert resp.json()["review_state"] == "draft"


def _complete_everything(client, session, aid):
    root = "Contoso tenant (demo)"
    for ref in ("IAM-01", "LOG-01"):
        as_user(client, PREPARER)
        answer(client, session, aid, ref, root, **COMPLETE)
        transition(client, session, aid, ref, root, "prepare")
        as_user(client, REVIEWER)
        transition(client, session, aid, ref, root, "review")


def test_close_requires_complete_reviewed_answers_then_locks(client, session, setup):
    aid = setup["aid"]
    as_user(client, REVIEWER)
    resp = client.post(f"/assessments/{aid}/close", data={"statement": "ok", "confirm": "on"},
                       follow_redirects=False)
    assert "Not ready" in unquote(resp.headers["location"])
    assert "missing or incomplete" in client.get(f"/assessments/{aid}/close").text

    _complete_everything(client, session, aid)
    resp = client.post(f"/assessments/{aid}/close", data={"statement": "Attested.", "confirm": "on"},
                       follow_redirects=False)
    assert "closed and attested" in unquote(resp.headers["location"])

    session.expire_all()
    r = session.scalar(select(Response))
    as_user(client, PREPARER)
    assert answer(client, session, aid, "IAM-01", "Contoso tenant (demo)", version=r.version,
                  narrative="after close").status_code == 423
    resp = client.post(f"/assessments/{aid}/artifacts", files={"file": ("late.txt", b"x", "text/plain")},
                       follow_redirects=False)
    assert "closed" in unquote(resp.headers["location"])
    resp = client.post(f"/assessments/{aid}/responses/comment", data={
        "item_id": item_id(session, aid, "IAM-01"), "node_id": node_id(session, aid), "body": "late"},
        follow_redirects=False)
    assert "closed" in unquote(resp.headers["location"])
    as_user(client, REVIEWER)
    assert client.post(f"/assessments/{aid}/reopen", data={"reason": "x"}).status_code == 403
    as_user(client, ADMIN)
    resp = client.post(f"/assessments/{aid}/reopen", data={"reason": ""}, follow_redirects=False)
    assert "reason is required" in unquote(resp.headers["location"])
    client.post(f"/assessments/{aid}/reopen", data={"reason": "auditor found a typo"})
    assert session.get(assessments.Assessment, aid).status == "open"


def test_issues_lifecycle(client, session, setup):
    aid = setup["aid"]
    iid = item_id(session, aid, "IAM-01")
    sub = node_id(session, aid, "sub-app-a-dev")
    client.post(f"/assessments/{aid}/issues", data={
        "item_id": iid, "node_id": sub, "title": "Break-glass account excluded from MFA",
        "severity": "high", "owner": "iam@corp.com", "due_date": "2026-11-30"})
    issue = session.scalar(select(Issue))
    assert issue.ref == "ISS-0001" and issue.scope_node_id == sub
    assert "Break-glass" in client.get("/issues").text

    base = {"title": issue.title, "severity": "high", "owner": "iam@corp.com", "due_date": "2026-11-30"}
    resp = client.post(f"/issues/{issue.id}", data=base | {"version": issue.version, "status": "closed"},
                       follow_redirects=False)
    assert "Say why" in unquote(resp.headers["location"])
    client.post(f"/issues/{issue.id}", data=base | {"version": issue.version, "status": "closed",
                                                    "note": "exclusion removed, see PR 42"})
    resp = client.post(f"/issues/{issue.id}", data=base | {"version": issue.version, "status": "open"},
                       follow_redirects=False)
    assert "Someone else changed" in unquote(resp.headers["location"])  # stale version
    session.expire_all()
    issue = session.get(Issue, issue.id)
    assert issue.status == "closed" and issue.closed_at is not None
    assert "exclusion removed" in client.get(f"/issues/{issue.id}").text


def test_excel_export(client, session, setup):
    aid = setup["aid"]
    answer(client, session, aid, "IAM-01", "mg-prod", **(COMPLETE | {
        "narrative": "=HYPERLINK(\"http://evil\",\"click\")"}))
    client.post(f"/assessments/{aid}/responses/evidence", data={
        "item_id": item_id(session, aid, "IAM-01"), "node_id": node_id(session, aid, "mg-prod"),
        "url": "https://github.example.com/org/audit/blob/main/mfa.md"})
    resp = client.get(f"/assessments/{aid}/export.xlsx")
    assert resp.status_code == 200
    wb = load_workbook(io.BytesIO(resp.content))
    assert wb.sheetnames == ["Summary", "Controls", "Effective answers", "Answers (all levels)",
                             "Issues", "Evidence", "Files", "Audit log"]
    eff = list(wb["Effective answers"].iter_rows(min_row=2, values_only=True))
    assert len(eff) == 2 * 5  # 2 controls x 5 subscriptions
    inherited = [row for row in eff if row[0] == "IAM-01" and row[7] == "inherited"]
    assert {row[3] for row in inherited} == {"sub-app-a-prod", "sub-app-b-prod"}
    assert "NOT pinned" in inherited[0][19]
    cell = wb["Answers (all levels)"]["H2"]
    assert cell.data_type == "s" and cell.value.startswith("=HYPERLINK")  # text, not a formula
    assert client.get(f"/assessments/{aid}/report").status_code == 200


def test_scope_sync(client, session, monkeypatch):
    from app import config
    monkeypatch.setattr(config, "AZURE_ROOT_MANAGEMENT_GROUP", "tenant-guid")
    root = {"azure_id": "tenant-guid", "kind": "tenant", "name": "Corp tenant", "parent": None}
    first = [  # child listed before its parent, as the API may do
        {"azure_id": "sub-1", "kind": "subscription", "name": "prod-1", "parent": "mg-prod"},
        {"azure_id": "mg-prod", "kind": "management_group", "name": "Prod", "parent": "tenant-guid"},
        {"azure_id": "sub-2", "kind": "subscription", "name": "prod-2", "parent": "mg-prod"},
    ]
    monkeypatch.setattr(scope, "fetch_hierarchy", lambda rid: (root, first))
    resp = client.post("/scope/sync", follow_redirects=False)
    assert "4 added" in unquote(resp.headers["location"])
    nodes = {n.azure_id: n for n in session.scalars(select(ScopeNode))}
    assert nodes["sub-1"].parent_id == nodes["mg-prod"].id
    assert nodes["mg-prod"].parent_id == nodes["tenant-guid"].id

    second = [first[1], dict(first[0], name="prod-1-renamed")]  # sub-2 gone
    monkeypatch.setattr(scope, "fetch_hierarchy", lambda rid: (root, second))
    resp = client.post("/scope/sync", follow_redirects=False)
    assert "1 updated, 1 deactivated" in unquote(resp.headers["location"])
    session.expire_all()
    assert session.get(ScopeNode, nodes["sub-2"].id).active is False
    assert session.get(ScopeNode, nodes["sub-1"].id).name == "prod-1-renamed"


def test_roles_on_setup_pages(client, session, setup):
    as_user(client, VIEWER)
    assert client.post("/libraries", data={"name": "x"}).status_code == 403
    assert client.post("/scope/sync").status_code == 403
    assert client.get("/assessments/new/controls").status_code == 403
    assert client.get(f"/assessments/{setup['aid']}").status_code == 200  # viewers can read
    as_user(client, PREPARER)
    assert client.get("/assessments/new/controls").status_code == 200


def test_overview_filters(client, session, setup):
    aid = setup["aid"]
    page = client.get(f"/assessments/{aid}?q=AU-11")
    assert "LOG-01" in page.text and "IAM-01" not in page.text
    page = client.get(f"/assessments/{aid}?category=Identity")
    assert "IAM-01" in page.text and "LOG-01" not in page.text


def test_every_page_renders(client, session, setup):
    aid = setup["aid"]
    iid = item_id(session, aid, "IAM-01")
    answer(client, session, aid, "IAM-01", "mg-prod", **COMPLETE)
    client.post(f"/assessments/{aid}/issues", data={"item_id": iid, "node_id": node_id(session, aid),
                                                    "title": "gap", "severity": "low"})
    lib_item = session.scalar(select(LibraryItem).where(LibraryItem.library_id == setup["lid"]))
    rid = session.scalar(select(Response.id))
    for url in ["/", "/libraries", f"/libraries/{setup['lid']}", f"/libraries/{setup['lid']}?show=all",
                f"/libraries/{setup['lid']}/items/{lib_item.id}", "/scope", "/assessments/new/controls",
                "/assessments/new/ai", f"/assessments/{aid}", f"/assessments/{aid}?show=gaps",
                f"/assessments/{aid}/items/{iid}", f"/assessments/{aid}/items/{iid}?node={node_id(session, aid, 'mg-prod')}",
                f"/assessments/{aid}/close", f"/assessments/{aid}/activity", f"/assessments/{aid}/issues",
                "/issues", "/issues?status=", "/issues/1", f"/assessments/{aid}/report", f"/responses/{rid}/history",
                "/admin/users", "/settings"]:
        resp = client.get(url)
        assert resp.status_code == 200, url
    assert client.get("/assessments/999", headers={"accept": "text/html"}).status_code == 404
